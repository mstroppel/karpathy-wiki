"""Private XLSX parser subprocess. Never print parser exceptions or warnings."""

from __future__ import annotations

import datetime
import io
import json
import re
import resource
import sys
import warnings
import zipfile
from decimal import Decimal
from pathlib import Path
from typing import Any

from karpathy_wiki_ingest_webdav.xlsx import (
    EXIT_CODES,
    EXTRACTION_VERSION,
    MAX_CELLS,
    MAX_COLUMNS,
    MAX_COMPRESSED,
    MAX_EXAMINED_ROWS,
    MAX_MERGES,
    MAX_OUTPUT,
    MAX_ROWS,
    MAX_SHEETS,
    MAX_UNCOMPRESSED,
    WorkbookError,
)


def preflight(source: Path) -> None:
    """Bound ZIP expansion and XML work before openpyxl allocates its object graph."""
    from defusedxml.ElementTree import iterparse
    from openpyxl.utils.cell import coordinate_to_tuple

    if source.stat().st_size > MAX_COMPRESSED:
        raise WorkbookError("limits")
    with zipfile.ZipFile(source) as archive:
        entries = archive.infolist()
        if len(entries) > 2048 or sum(item.file_size for item in entries) > MAX_UNCOMPRESSED:
            raise WorkbookError("limits")
        names = [item.filename for item in entries]
        if len(set(names)) != len(names):
            raise WorkbookError("invalid")
        cells = rows = merges = elements = sheets = 0
        workbook_type = False
        for entry in entries:
            if entry.flag_bits & 1:
                raise WorkbookError("invalid")
            if entry.compress_type not in {zipfile.ZIP_STORED, zipfile.ZIP_DEFLATED}:
                raise WorkbookError("unsupported")
            if (
                entry.file_size > max(entry.compress_size, 1) * 200
                and entry.file_size > 1024 * 1024
            ):
                raise WorkbookError("limits")
            if "vbaproject" in entry.filename.lower():
                raise WorkbookError("unsupported")
            # Read every member once for CRC/actual expansion validation. No extraction,
            # relationship traversal, external entity resolution or network access.
            with archive.open(entry) as handle:
                data = handle.read(MAX_UNCOMPRESSED + 1)
            if len(data) != entry.file_size or len(data) > MAX_UNCOMPRESSED:
                raise WorkbookError("limits")
            if not entry.filename.endswith((".xml", ".rels")):
                continue
            for _, element in iterparse(io.BytesIO(data), events=("end",), forbid_dtd=True):
                elements += 1
                if elements > 1_000_000:
                    raise WorkbookError("limits")
                tag = element.tag.rsplit("}", 1)[-1]
                if tag == "Override" and "macroEnabled" in element.get("ContentType", ""):
                    raise WorkbookError("unsupported")
                if entry.filename == "[Content_Types].xml" and tag == "Override":
                    if element.get("ContentType") == (
                        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml"
                    ):
                        workbook_type = True
                spreadsheet_tag = element.tag.startswith(
                    "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
                )
                if spreadsheet_tag and tag == "sheet":
                    sheets += 1
                    if sheets > MAX_SHEETS:
                        raise WorkbookError("limits")
                if spreadsheet_tag:
                    if tag == "row":
                        rows += 1
                        if rows > MAX_EXAMINED_ROWS or int(element.get("r", "0")) > MAX_ROWS:
                            raise WorkbookError("limits")
                    elif tag == "c":
                        cells += 1
                        row, column = coordinate_to_tuple(element.get("r", ""))
                        if (
                            cells > MAX_CELLS
                            or not 1 <= row <= MAX_ROWS
                            or not 1 <= column <= MAX_COLUMNS
                        ):
                            raise WorkbookError("limits")
                    elif tag == "mergeCell":
                        merges += 1
                        if merges > MAX_MERGES:
                            raise WorkbookError("limits")
                        from openpyxl.worksheet.cell_range import CellRange

                        area = CellRange(element.get("ref", ""))
                        if (
                            area.max_row > MAX_ROWS
                            or area.max_col > MAX_COLUMNS
                            or area.size["rows"] * area.size["columns"] > MAX_CELLS
                        ):
                            raise WorkbookError("limits")
                element.clear()
        if not workbook_type:
            raise WorkbookError("unsupported")


def value_text(cell: Any) -> str:
    value = cell.value
    if value is None:
        return ""
    if cell.data_type == "e":
        return f"Excel error: {value}"
    if isinstance(value, bool):
        return f"boolean: {str(value).lower()}"
    if isinstance(value, (datetime.datetime, datetime.date, datetime.time)):
        return f"date/time: {value.isoformat()} (Excel format: {cell.number_format})"
    if isinstance(value, datetime.timedelta):
        return f"duration: {value.total_seconds()} seconds (Excel format: {cell.number_format})"
    if not isinstance(value, (int, float)):
        return str(value)
    fmt = cell.number_format
    if fmt == "General":
        return f"number: {value}"
    # Deliberately small, documented subset; never pretend to render all of Excel.
    match = re.fullmatch(r'(0|#,##0)(\.0{1,6})?(%| "[^"\r\n]+")?', fmt)
    if match:
        suffix = match[3] or ""
        places = len(match[2]) - 1 if match[2] else 0
        number = Decimal(str(value)) * (100 if suffix == "%" else 1)
        rendered = f"{number:.{places}f}"
        if suffix == "%":
            return f"percentage: {rendered}%"
        unit = f" (unit: {suffix[2:-1]})" if suffix else ""
        return f"number: {rendered}{unit}"
    return f"number: {value} (unsupported number format: {fmt})"


def segments(source: Path) -> list[list[Any]]:
    from openpyxl import load_workbook
    from openpyxl.cell.cell import MergedCell
    from openpyxl.utils.cell import get_column_letter

    preflight(source)
    formulas = load_workbook(source, data_only=False, keep_links=False)
    cached = load_workbook(source, data_only=True, keep_links=False)
    result: list[list[Any]] = []
    size = examined = 0

    def add(text: str, *, user: bool = False) -> None:
        nonlocal size
        segment = [user, text]
        size += len(json.dumps(segment, ensure_ascii=False).encode("utf-8")) + 2
        if size > MAX_OUTPUT - 2:
            raise WorkbookError("limits")
        result.append(segment)

    add(
        f"Extraction version: {EXTRACTION_VERSION}\n\n"
        "Visible worksheets, rows and columns only. "
        "Hidden/very-hidden sheets, hidden rows/columns, "
        "comments, document properties, charts, images and link targets are omitted. "
        "Empty outer ranges are omitted; interior blanks retain coordinates. "
        "Formatting is a limited semantic extract, not Excel rendering.\n\n"
        "Formulas are calculation context, not executed or verified. Cached results may be stale; "
        "recalculate and save in Excel before upload. "
        "Missing caches and Excel errors are explicit.\n"
    )
    try:
        if len(formulas.worksheets) > MAX_SHEETS:
            raise WorkbookError("limits")
        for sheet in formulas.worksheets:
            if sheet.sheet_state != "visible":
                continue
            # openpyxl does not retain whether hidden="0" was explicitly set on
            # individual rows. Never guess visibility under a default-hidden sheet.
            if (
                sheet.sheet_format.zeroHeight
                or sheet.sheet_format.defaultRowHeight == 0
                or sheet.sheet_format.defaultColWidth == 0
            ):
                raise WorkbookError("visibility")
            add("\n## Worksheet: ")
            add(sheet.title, user=True)
            add("\n\n")
            hidden_columns: set[int] = set()
            for dimension in sheet.column_dimensions.values():
                if dimension.hidden or dimension.width == 0:
                    hidden_columns.update(
                        range(dimension.min or 1, min(dimension.max or 1, MAX_COLUMNS) + 1)
                    )

            def visible(row: int, column: int, hidden_columns=hidden_columns, sheet=sheet) -> bool:
                dimension = sheet.row_dimensions[row]
                return (
                    column not in hidden_columns and not dimension.hidden and dimension.height != 0
                )

            # Normal mode retains merged/visibility metadata; ZIP/XML preflight plus
            # OS limits bound its memory. Ignore misleading <dimension> declarations.
            bounds: list[tuple[int, int]] = []
            for cell in sheet._cells.values():
                if cell.value is not None and visible(cell.row, cell.column):
                    bounds.append((cell.row, cell.column))
            if not bounds:
                add("No visible non-empty cells.\n")
                continue
            for area in sheet.merged_cells.ranges:
                if sheet.cell(area.min_row, area.min_col).value is not None and visible(
                    area.min_row, area.min_col
                ):
                    bounds.append((area.max_row, area.max_col))
            min_row = min(row for row, _ in bounds)
            max_row = max(row for row, _ in bounds)
            min_col = min(column for _, column in bounds)
            max_col = max(column for _, column in bounds)
            examined += (max_row - min_row + 1) * (max_col - min_col + 1)
            if max_row > MAX_ROWS or max_col > MAX_COLUMNS or examined > MAX_CELLS:
                raise WorkbookError("limits")
            add(
                f"Range: {get_column_letter(min_col)}{min_row}:"
                f"{get_column_letter(max_col)}{max_row}\n\n"
            )
            for area in sorted(sheet.merged_cells.ranges, key=str):
                if any(
                    min_row <= row <= max_row and min_col <= col <= max_col and visible(row, col)
                    for row, col in ((area.min_row, area.min_col), (area.max_row, area.max_col))
                ):
                    add(
                        f"Merged range: {area}; value only at anchor "
                        f"{area.start_cell.coordinate}.\n"
                    )
            columns = [col for col in range(min_col, max_col + 1) if col not in hidden_columns]
            add("\n| Row | " + " | ".join(get_column_letter(col) for col in columns) + " |\n")
            add("| --- | " + " | ".join("---" for _ in columns) + " |\n")
            for row in range(min_row, max_row + 1):
                if not visible(row, columns[0]):
                    continue
                add(f"| {row} | ")
                for col in columns:
                    cell = sheet.cell(row, col)
                    if isinstance(cell, MergedCell):
                        text = "merged (no independent value)"
                    elif cell.data_type == "f":
                        if not isinstance(cell.value, str):
                            # Array/data-table formula objects are not simple formula strings.
                            raise WorkbookError("formula")
                        cache = cached[sheet.title].cell(row, col)
                        value = (
                            value_text(cache)
                            if cache.value is not None
                            else "MISSING cached result"
                        )
                        text = f"formula: {cell.value}; cached: {value}"
                    else:
                        text = value_text(cell)
                    add(text, user=True)
                    add(" | ")
                add("\n")
    finally:
        formulas.close()
        cached.close()
    return result


def main() -> None:
    # Applied before importing openpyxl/defusedxml, independently for each workbook.
    resource.setrlimit(resource.RLIMIT_AS, (768 * 1024 * 1024,) * 2)
    resource.setrlimit(resource.RLIMIT_CPU, (30, 30))
    resource.setrlimit(resource.RLIMIT_FSIZE, (MAX_OUTPUT, MAX_OUTPUT))
    resource.setrlimit(resource.RLIMIT_CORE, (0, 0))
    warnings.simplefilter("ignore")
    try:
        output = segments(Path(sys.argv[1]))
        sys.stdout.write(json.dumps(output, ensure_ascii=False, separators=(",", ":")))
    except WorkbookError as error:
        raise SystemExit(EXIT_CODES[error.code]) from None
    except MemoryError:
        raise SystemExit(EXIT_CODES["resources"]) from None
    except Exception:
        # No raw parser errors, workbook strings or traceback in logs/reports.
        raise SystemExit(2) from None


if __name__ == "__main__":
    main()
