"""Synthetic workbooks only: extraction, privacy and generation regression coverage."""

import datetime
import json
import subprocess
import tempfile
import unittest
import zipfile
from contextlib import closing
from pathlib import Path
from unittest import mock
from xml.etree import ElementTree as ET

from defusedxml.common import DefusedXmlException
from openpyxl import Workbook
from openpyxl.comments import Comment

import karpathy_wiki_ingest_webdav as webdav
from karpathy_wiki_ingest.shared import TargetedAnonymizer
from karpathy_wiki_ingest.state import StateStore
from karpathy_wiki_ingest_webdav import xlsx, xlsx_worker

NS = {"s": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}


def rewrite(path, transform):
    with zipfile.ZipFile(path) as archive:
        members = {name: archive.read(name) for name in archive.namelist()}
    transform(members)
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as archive:
        for name, data in members.items():
            archive.writestr(name, data)


def caches(path):
    def transform(members):
        root = ET.fromstring(members["xl/worksheets/sheet1.xml"])
        for cell in root.findall(".//s:c", NS):
            if cell.get("r") in {"C3", "D3"}:
                value = cell.find("s:v", NS)
                assert value is not None
                value.text = "3" if cell.get("r") == "C3" else "#DIV/0!"
                if cell.get("r") == "D3":
                    cell.set("t", "e")
        members["xl/worksheets/sheet1.xml"] = ET.tostring(root)

    rewrite(path, transform)


class WorkbookFixture(unittest.TestCase):
    def setUp(self):
        self.temporary = tempfile.TemporaryDirectory()
        self.addCleanup(self.temporary.cleanup)
        self.root = Path(self.temporary.name)
        self.source = self.root / "Max Mustermann.xlsx"
        self.relative = Path("Max Mustermann.xlsx")
        self.anonymizer = TargetedAnonymizer.from_config(
            {"people": [{"values": ["Max Mustermann"], "replacement": "[PERSON_1]"}]}
        )

    def workbook(self):
        book = Workbook()
        book.active.title = "Max Mustermann"
        return book

    def render(self):
        # Exercise the parser in-process for branch coverage, with the same protocol
        # used by the resource-limited subprocess in the integration tests below.
        return xlsx.render(xlsx_worker.segments(self.source), self.relative, self.anonymizer)


class WorkbookTests(WorkbookFixture):
    def test_multisheet_unicode_blanks_headers_merged_and_escaping(self):
        book = self.workbook()
        sheet = book.active
        sheet.append(["Header α", "Second", "Third"])
        sheet.append(["雪 | <b> & [link](https://example.invalid) `code`", None, "end\nline"])
        sheet.append(["Max Mustermann", "", "last"])
        sheet.merge_cells("A4:C4")
        sheet["A4"] = "merged text"
        sheet["A5"] = "last row"
        sheet["Z100"].number_format = "0.00"  # Empty formatting must not expand the extract.
        second = book.create_sheet("Résumé")
        second["C3"] = "value"
        book.save(self.source)
        output = self.render()
        self.assertEqual(output, self.render())
        self.assertIn("Worksheet: \\[PERSON\\_1\\]", output)
        self.assertNotIn("Max Mustermann", output)
        self.assertIn("Range: A1:C5", output)
        self.assertIn("| 1 | Header α | Second | Third |", output)
        self.assertIn("| 2 | 雪 \\| &lt;b&gt; &amp; \\[link\\]", output)
        self.assertIn(" |  | end<br>line |", output)
        self.assertIn("Merged range: A4:C4; value only at anchor A4", output)
        self.assertIn("merged \\(no independent value\\)", output)
        self.assertIn("Worksheet: Résumé", output)
        self.assertIn("Range: C3:C3", output)
        self.assertNotIn("Z100", output)

    def test_hidden_sheets_rows_grouped_columns_and_comments_are_omitted(self):
        book = self.workbook()
        sheet = book.active
        sheet.append(["visible", "HIDDEN-COL", "HIDDEN-GROUP", "HIDDEN-GROUP", "end"])
        sheet.append(["HIDDEN-ROW"])
        sheet.column_dimensions.group("B", "D", hidden=True)
        sheet.row_dimensions[2].hidden = True
        sheet["A1"].comment = Comment("COMMENT-SECRET", "AUTHOR-SECRET")
        for state in ("hidden", "veryHidden"):
            hidden = book.create_sheet(f"SHEET-SECRET-{state}")
            hidden.sheet_state = state
            hidden["A1"] = "HIDDEN-VALUE"
        book.properties.title = "PROPERTY-SECRET"
        book.save(self.source)
        output = self.render()
        self.assertIn("| Row | A | E |", output)
        for secret in ("HIDDEN-", "SECRET", "| 2 |"):
            self.assertNotIn(secret, output)
        self.assertIn("Hidden/very-hidden", output)

    def test_dates_percentages_units_booleans_numbers_and_unsupported_formats(self):
        book = self.workbook()
        sheet = book.active
        sheet["A1"] = datetime.date(2026, 10, 22)
        sheet["A1"].number_format = "yyyy-mm-dd"
        sheet["B1"] = 0.125
        sheet["B1"].number_format = "0.0%"
        sheet["C1"] = 12.5
        sheet["C1"].number_format = '0.00 "kg"'
        sheet["D1"] = True
        sheet["E1"] = 1234.567
        sheet["E1"].number_format = "#,##0.00"
        sheet["F1"] = -2
        sheet["F1"].number_format = "[Red]0;0"
        sheet["G1"] = "#N/A"
        sheet["H1"] = 0
        sheet["I1"] = datetime.time(12, 30)
        sheet["J1"] = datetime.timedelta(hours=30)
        sheet["J1"].number_format = "[h]:mm:ss"
        book.save(self.source)
        output = self.render()
        for text in (
            "date/time: 2026\\-10\\-22T00:00:00",
            "percentage: 12\\.5%",
            "number: 12\\.50 \\(unit: kg\\)",
            "boolean: true",
            "number: 1234\\.57",
            "unsupported number format",
            "Excel error: \\#N/A",
            "number: 0",
            "date/time: 12:30:00",
            "duration: 108000\\.0 seconds",
        ):
            self.assertIn(text, output)

    def test_formula_caches_missing_errors_and_redaction_without_calculation(self):
        book = self.workbook()
        sheet = book.active
        sheet["A1"] = 1
        sheet["B1"] = 2
        sheet["A3"] = '=IF(A1=1,"Max Mustermann","")'
        sheet["B3"] = "=A1+B1"
        sheet["C3"] = "=A1+B1"
        sheet["D3"] = "=1/0"
        book.save(self.source)
        caches(self.source)
        output = self.render()
        self.assertIn("formula: =A1\\+B1; cached: MISSING cached result", output)
        self.assertIn("formula: =A1\\+B1; cached: number: 3", output)
        self.assertIn("formula: =1/0; cached: Excel error: \\#DIV/0\\!", output)
        self.assertNotIn("Mustermann", output)
        self.assertIn("recalculate and save in Excel", output)
        self.assertIn("not executed or verified", output)

    def test_raw_values_are_redacted_before_markdown_escaping(self):
        book = self.workbook()
        book.active["A1"] = "Mary*Smith & Sons"
        book.active["B1"] = '=IF(A1="Mary*Smith & Sons",1,0)'
        book.save(self.source)
        redactor = TargetedAnonymizer.from_config(
            {"people": [{"values": ["Mary*Smith & Sons"], "replacement": "[PERSON]"}]}
        )
        output = xlsx.render(xlsx_worker.segments(self.source), self.relative, redactor)
        self.assertNotIn("Smith", output)
        self.assertIn("\\[PERSON\\]", output)

    def test_empty_visible_sheet_is_reported_and_hidden_anchor_does_not_leak(self):
        book = self.workbook()
        book.active.merge_cells("A1:C1")
        book.active["A1"] = "HIDDEN-ANCHOR"
        book.active.column_dimensions["A"].hidden = True
        book.save(self.source)
        self.assertIn("No visible non-empty cells", self.render())
        self.assertNotIn("HIDDEN-ANCHOR", self.render())

    def test_default_hidden_rows_fail_closed_instead_of_guessing_visibility(self):
        book = self.workbook()
        book.active["A1"] = "DEFAULT-HIDDEN-SECRET"
        book.active.sheet_format.zeroHeight = True
        book.save(self.source)
        with self.assertRaisesRegex(xlsx.WorkbookError, "copy visible data"):
            self.render()
        with self.assertRaisesRegex(xlsx.WorkbookError, "copy visible data"):
            xlsx.extract(self.source, self.relative, self.anonymizer)

    def test_no_external_fetch_and_unsupported_array_formula_fails(self):
        book = self.workbook()
        book.active["A1"] = "='[https://example.invalid/private.xlsx]Sheet1'!A1"
        book.active["A2"] = "safe label"
        book.active["A2"].hyperlink = "https://example.invalid/LINK-TARGET-SECRET"
        book.save(self.source)
        with mock.patch("urllib.request.urlopen", side_effect=AssertionError("network access")):
            output = self.render()
        self.assertIn("MISSING cached result", output)
        self.assertNotIn("LINK-TARGET-SECRET", output)
        from openpyxl.worksheet.formula import ArrayFormula

        book.active["A1"] = ArrayFormula(ref="A1:A2", text="=SUM(A2)")
        book.save(self.source)
        with self.assertRaisesRegex(xlsx.WorkbookError, "array/data-table formula"):
            self.render()

    def test_extraction_limits_fail_explicitly_not_by_truncation(self):
        book = self.workbook()
        book.active["A1"] = "safe"
        book.active["C3"] = "more"
        book.active.merge_cells("A4:C4")
        book.create_sheet("Second")["A1"] = "other"
        book.save(self.source)
        for constant, limit in (
            ("MAX_COMPRESSED", 1),
            ("MAX_UNCOMPRESSED", 1),
            ("MAX_SHEETS", 1),
            ("MAX_ROWS", 1),
            ("MAX_COLUMNS", 1),
            ("MAX_EXAMINED_ROWS", 1),
            ("MAX_CELLS", 1),
            ("MAX_MERGES", 0),
            ("MAX_OUTPUT", 32),
        ):
            with self.subTest(constant=constant), mock.patch.object(xlsx_worker, constant, limit):
                with self.assertRaisesRegex(xlsx.WorkbookError, "exceeds extraction limits"):
                    self.render()

    def test_sparse_rectangle_and_large_merge_are_bounded(self):
        book = self.workbook()
        book.active["A1"] = "first"
        book.active["IV10000"] = "last"
        book.save(self.source)
        with self.assertRaisesRegex(xlsx.WorkbookError, "exceeds extraction limits"):
            self.render()

        def add_large_merge(members):
            members["xl/worksheets/sheet1.xml"] = members["xl/worksheets/sheet1.xml"].replace(
                b"</worksheet>",
                b'<mergeCells><mergeCell ref="A1:IV10000"/></mergeCells></worksheet>',
            )

        rewrite(self.source, add_large_merge)
        with self.assertRaisesRegex(xlsx.WorkbookError, "exceeds extraction limits"):
            self.render()

    def test_zip_bombs_duplicates_macros_and_dtds_are_rejected(self):
        for kind in ("ratio", "entries", "duplicate", "macro", "dtd"):
            with self.subTest(kind=kind):
                book = self.workbook()
                book.save(self.source)

                def transform(members, kind=kind):
                    if kind == "ratio":
                        members["bomb.bin"] = b"x" * (2 * 1024 * 1024)
                    elif kind == "entries":
                        members.update({f"entry-{i}": b"" for i in range(2049)})
                    elif kind == "macro":
                        members["xl/vbaProject.bin"] = b"synthetic macro placeholder"
                    elif kind == "dtd":
                        members["xl/worksheets/sheet1.xml"] = (
                            b'<!DOCTYPE sheet [<!ENTITY x "PRIVATE-PARSER-ERROR">]>'
                            b"<sheet>&x;</sheet>"
                        )

                rewrite(self.source, transform)
                if kind == "duplicate":
                    with zipfile.ZipFile(self.source, "a") as archive:
                        archive.writestr("xl/workbook.xml", b"duplicate")
                with self.assertRaises((xlsx.WorkbookError, DefusedXmlException)):
                    self.render()

    def test_subprocess_success_corrupt_unsupported_and_safe_resource_failures(self):
        book = self.workbook()
        book.active["A1"] = "Max Mustermann"
        book.save(self.source)
        self.assertEqual(xlsx.extract(self.source, self.relative, self.anonymizer), self.render())
        for suffix in (".xls", ".xlsb", ".xlsm"):
            with self.assertRaisesRegex(xlsx.WorkbookError, "unsupported"):
                xlsx.extract(self.source.with_suffix(suffix), self.relative, self.anonymizer)
        self.source.write_bytes(b"PRIVATE-PARSER-ERROR\x00corrupt")
        with self.assertRaisesRegex(xlsx.WorkbookError, "Invalid or encrypted"):
            xlsx.extract(self.source, self.relative, self.anonymizer)
        with mock.patch.object(xlsx, "MAX_COMPRESSED", 1):
            with self.assertRaisesRegex(xlsx.WorkbookError, "exceeds extraction limits"):
                xlsx.extract(self.source, self.relative, self.anonymizer)
        for outcome in (-9, 3, 4, 2):
            with mock.patch.object(
                xlsx.subprocess, "run", return_value=mock.Mock(returncode=outcome)
            ):
                with self.assertRaises(xlsx.WorkbookError):
                    xlsx.extract(self.source, self.relative, self.anonymizer)
        with mock.patch.object(
            xlsx.subprocess, "run", side_effect=subprocess.TimeoutExpired("private", 1)
        ):
            with self.assertRaisesRegex(xlsx.WorkbookError, "parser resources"):
                xlsx.extract(self.source, self.relative, self.anonymizer)

    def test_invalid_protocol_and_escaped_output_size_fail_closed(self):
        for segments in ([[0, "bad"]], [[True]], [[True, 1]], ["bad"]):
            with self.assertRaises(xlsx.WorkbookError):
                xlsx.render(segments, self.relative, self.anonymizer)
        with mock.patch.object(xlsx, "MAX_OUTPUT", 100):
            with self.assertRaisesRegex(xlsx.WorkbookError, "exceeds extraction limits"):
                xlsx.render([[True, "<" * 100]], self.relative, self.anonymizer)

    def test_worker_applies_os_limits_and_returns_only_fixed_failure_codes(self):
        for error, code in (
            (MemoryError("private"), 5),
            (ValueError("private"), 2),
            (xlsx.WorkbookError("limits"), 3),
        ):
            with (
                mock.patch.object(xlsx_worker.resource, "setrlimit") as limits,
                mock.patch.object(xlsx_worker, "segments", side_effect=error),
                mock.patch.object(xlsx_worker.sys, "argv", ["worker", str(self.source)]),
                self.assertRaises(SystemExit) as raised,
            ):
                xlsx_worker.main()
            self.assertEqual(raised.exception.code, code)
            self.assertEqual(limits.call_count, 4)
            limits.assert_any_call(xlsx_worker.resource.RLIMIT_CPU, (30, 30))
            limits.assert_any_call(xlsx_worker.resource.RLIMIT_FSIZE, (xlsx.MAX_OUTPUT,) * 2)


class GenerationTests(WorkbookFixture):
    def test_oversize_inventory_does_not_hash_the_rejected_input(self):
        book = self.workbook()
        book.save(self.source)
        with (
            mock.patch.object(xlsx, "MAX_COMPRESSED", 1),
            mock.patch.object(
                webdav, "file_revision", side_effect=AssertionError("unbounded read")
            ),
        ):
            inventory = webdav.upstream_inventory(self.root)
        self.assertEqual(len(inventory), 1)
        self.assertNotIn("Mustermann", json.dumps(inventory))

    def test_generation_lifecycle_privacy_revision_redaction_and_extraction_changes(self):
        incoming = self.root / "incoming"
        incoming.mkdir()
        source = incoming / self.relative
        sanitized, quarantine = self.root / "published", self.root / "quarantine"
        book = self.workbook()
        book.active["A1"] = "Max Mustermann"
        book.save(source)
        (incoming / "Max Mustermann.xlsx.md").write_text("ordinary Markdown")
        with closing(StateStore.open(self.root / "state.sqlite3")) as store:

            def cycle(redactor=self.anonymizer):
                return webdav.process_cycle(
                    incoming, sanitized, quarantine, redactor, store, interval=60
                )

            self.assertEqual(cycle(), (2, 0, False))
            # Ordinary Markdown keeps its existing path behavior; remove it before
            # verifying the stricter workbook metadata boundary.
            (incoming / "Max Mustermann.xlsx.md").unlink()
            self.assertEqual(cycle(), (0, 0, False))
            manifest = json.loads((sanitized / "manifest.json").read_text())
            item = manifest["items"][0]
            key = xlsx.output_path(self.relative).as_posix()
            self.assertEqual(item["source_key"], key)
            self.assertEqual(item["wiki_path"], f"webdav/{key}/index.md")
            self.assertTrue(item["source_path"].endswith(f"/{key}"))
            for path in sanitized.rglob("*"):
                if path.is_file():
                    self.assertNotIn("Max Mustermann", path.read_text())
            previous = webdav.active_generation(sanitized)
            self.assertEqual(cycle(), (0, 0, False))
            self.assertEqual(webdav.active_generation(sanitized), previous)
            book.active["B1"] = "changed"
            book.save(source)
            self.assertEqual(cycle(), (1, 0, False))
            self.assertNotEqual(
                item["source_revision"],
                json.loads((sanitized / "manifest.json").read_text())["items"][0][
                    "source_revision"
                ],
            )
            rotated = TargetedAnonymizer.from_config(
                {"people": [{"values": ["Max Mustermann"], "replacement": "[AUTHOR]"}]}
            )
            self.assertEqual(cycle(rotated), (1, 0, False))
            previous = webdav.active_generation(sanitized)
            with mock.patch.object(webdav, "EXTRACTION_VERSION", 999):
                self.assertEqual(cycle(rotated), (0, 0, False))
                self.assertNotEqual(webdav.active_generation(sanitized), previous)
            source.unlink()
            self.assertEqual(cycle(rotated), (0, 0, False))
            self.assertEqual(json.loads((sanitized / "manifest.json").read_text())["items"], [])

    def test_failure_isolation_safe_reports_and_output_collisions(self):
        incoming = self.root / "incoming"
        incoming.mkdir()
        sanitized, quarantine = self.root / "published", self.root / "quarantine"
        (incoming / "ordinary.md").write_text("previous")

        def publish():
            return webdav.publish_generation(incoming, sanitized, quarantine, self.anonymizer)

        self.assertEqual(publish(), (1, 0))
        previous = webdav.active_generation(sanitized)
        manifest = (sanitized / "manifest.json").read_bytes()
        source = incoming / self.relative
        for contents in (b"PRIVATE-PARSER-ERROR", b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1encrypted"):
            source.write_bytes(contents)
            with self.assertLogs(webdav.LOG, level="ERROR") as logs:
                self.assertEqual(publish(), (0, 1))
            self.assertNotIn("Mustermann", "".join(logs.output))
            self.assertEqual(webdav.active_generation(sanitized), previous)
            self.assertEqual((sanitized / "manifest.json").read_bytes(), manifest)
            report = next(quarantine.rglob("*.error")).read_text()
            self.assertIn("unencrypted .xlsx", report)
            self.assertNotIn("Mustermann", report)
            self.assertNotIn("PRIVATE-PARSER-ERROR", report)
        book = self.workbook()
        book.active["A1"] = "safe"
        book.save(source)
        collision = incoming / xlsx.output_path(self.relative)
        collision.parent.mkdir(parents=True)
        collision.write_text("do not overwrite")
        self.assertEqual(publish(), (0, 1))
        self.assertEqual(webdav.active_generation(sanitized), previous)
        collision.unlink()
        self.assertEqual(publish(), (1, 0))


if __name__ == "__main__":
    unittest.main()
